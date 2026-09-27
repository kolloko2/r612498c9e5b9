"""Validate a teacher-maintained XLSX directory and export local JSON rules.

Columns: id, approved (yes/no), source, city, district (округ), area (район),
object, street, house, building, structure, services (one per line).
Does not publish, replace a directory, or invent geography.
"""
import argparse
import json
import os
import sys
from pathlib import Path
import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'backend'))
from territories import FIELDS, rules


def convert(source, output):
    if output.exists():
        raise ValueError('Выходной файл уже существует; выберите новый путь для проверки изменений')
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True, keep_links=False)
    try:
        sheet = workbook.active
        if sheet.max_row > 2001 or sheet.max_column > 20:
            raise ValueError('Не более 2000 правил и 20 столбцов')
        rows = sheet.iter_rows(values_only=True)
        headers = [str(v or '').strip() for v in next(rows)]
        if len(set(headers)) != len(headers) or not {'id','approved','source','services'} <= set(headers):
            raise ValueError('Нужны уникальные заголовки id, approved, source, services и адресные поля')
        if set(headers) - (FIELDS | {'id','approved','source','services'}):
            raise ValueError('Неизвестные столбцы')
        result = []
        for row in rows:
            if not any(v is not None for v in row):
                continue
            item = {k:str(v).strip() if v is not None else '' for k,v in zip(headers,row)}
            approved = item['approved'].casefold()
            if approved not in {'yes','no','да','нет','true','false'}:
                raise ValueError('approved: да/нет, yes/no или true/false')
            result.append({'id':item['id'], 'source':item['source'],
                           'approved':approved in {'yes','да','true'},
                           'match':{k:v for k,v in item.items() if k in FIELDS and v},
                           'services':[v.strip() for v in item['services'].splitlines() if v.strip()]})
    finally:
        workbook.close()
    # Reuse the runtime validator, with a task-local candidate file.
    candidate = output.with_suffix(output.suffix+'.candidate')
    created = False
    old = os.environ.get('TERRITORIAL_ROUTES_FILE')
    try:
        with candidate.open('x', encoding='utf-8') as stream:
            json.dump({'version':1,'rules':result},stream,ensure_ascii=False,indent=2)
        created = True
        os.environ['TERRITORIAL_ROUTES_FILE'] = str(candidate)
        active = rules()
        candidate.rename(output)
        return {'rules':len(result),'approved':len(active),'output':str(output)}
    finally:
        if old is None:
            os.environ.pop('TERRITORIAL_ROUTES_FILE',None)
        else:
            os.environ['TERRITORIAL_ROUTES_FILE'] = old
        if created and candidate.exists():
            candidate.unlink()


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source',type=Path)
    parser.add_argument('output',type=Path)
    args = parser.parse_args()
    print(json.dumps(convert(args.source,args.output),ensure_ascii=False))
