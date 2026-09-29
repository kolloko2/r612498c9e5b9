"""Build the editable training-library and coverage DOCX from project sources."""

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "output" / "documentation-2026-09-28"
MATERIALS = json.loads((ROOT / "tools" / "data" / "materials.json").read_text(encoding="utf-8"))


def shade(cell, color):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tcpr.append(shd)


def border(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    edges = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), "D9D9D9")
        edges.append(edge)
    tcpr.append(edges)


def base():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2)
    sec.left_margin, sec.right_margin = Cm(2.2), Cm(1.8)
    for name, size, before, after in (("Normal", 10.5, 0, 6), ("Title", 19, 0, 15),
                                      ("Heading 1", 14, 16, 7), ("Heading 2", 11, 10, 4)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.16
    return doc


def lead(doc, title, intro):
    doc.add_paragraph(title, "Title")
    doc.add_paragraph(intro)
    p = doc.add_paragraph("28.09.2026")
    p.style.font.size = Pt(9)
    p.paragraph_format.space_after = Pt(14)


def material_paragraphs(doc, body):
    for block in body.split("\n\n"):
        lines = block.splitlines()
        if len(lines) > 1 and lines[0].isupper() and len(lines[0]) < 70:
            doc.add_paragraph(lines[0].capitalize(), "Heading 2")
            lines = lines[1:]
        for line in lines:
            line = line.strip()
            if not line:
                continue
            if line[0].isdigit() and ". " in line[:4]:
                doc.add_paragraph(line, style="List Number")
            elif line.startswith("— "):
                doc.add_paragraph(line[2:], style="List Bullet")
            else:
                doc.add_paragraph(line)


def make_materials():
    doc = base()
    lead(doc, "Учебные материалы диспетчера ДДС",
         "Шесть коротких материалов помогают пройти полный цикл работы с учебной карточкой: "
         "принять её, передать задачу, отразить доклады, проверить расхождения и записать результат. "
         "Читайте материал перед соответствующим упражнением, затем сверяйте свои действия с разбором попытки.")
    doc.add_paragraph("Как пользоваться", "Heading 1")
    doc.add_paragraph("Начните с памятки о карточке и телефонном докладе. После вводного занятия "
                      "переходите к проверке ошибки 112 и новой вводной. В упражнениях действуют факты конкретного "
                      "сценария: примеры ниже не подставляются в карточку автоматически.")
    for i, item in enumerate(MATERIALS, 1):
        h = doc.add_paragraph(f"{i}. {item['title']}", "Heading 1")
        h.paragraph_format.keep_with_next = True
        doc.add_paragraph(item["description"])
        material_paragraphs(doc, item["body"])
    dest = OUT / "13_Учебные_материалы_ДДС.docx"
    doc.save(dest)
    return dest


ROWS = [
    ("Приём карточки", "Памятка диспетчера ДДС", "Открыть, принять или мотивированно отклонить", "Время открытия, первый статус и комментарий"),
    ("Передача задачи", "Доклад по телефону; Телефонный доклад", "Позвонить бригаде и начальнику смены", "История звонков, обязательные факты"),
    ("Ход реагирования", "Образцы комментариев; Телефонный доклад", "Записать доклад и сменить статус после него", "Последовательность статусов и содержание записей"),
    ("Ошибка карточки 112", "Проверка карточки 112", "Уточнить значение и сообщить в 112", "Исходное поле, уточнение, сообщение"),
    ("Новая вводная с пострадавшим", "Новая вводная: пострадавший и другая служба", "На специально заданном сценарии уточнить факт и передать его в 112", "Доклад, комментарий, сообщение в 112"),
    ("Завершение работы", "Памятка диспетчера; Образцы комментариев", "Записать результат и завершить карточку", "Итоговый статус и разбор попытки"),
]


def make_coverage():
    doc = base()
    lead(doc, "Покрытие учебных навыков и требований",
         "Матрица показывает, чему учит каждый материал и какое действие в тренажёре подтверждает навык. "
         "Наличие памятки само по себе не означает успешного прохождения: вывод делают по сохранённой попытке ученика.")
    doc.add_paragraph("Порядок занятия", "Heading 1")
    for step in (
        "Прочитать памятку о карточке и телефонном докладе.",
        "Пройти вводное занятие с подсказками: после доклада выбрать статус и записать услышанное.",
        "Повторить карточку без подсказок, проверить адрес и исправление ошибки 112.",
        "Для новой вводной использовать сценарий, где преподаватель явно задал пострадавшего.",
        "Разобрать историю звонков, статусов, комментариев и сообщений в 112 вместе с оценкой.",
    ):
        doc.add_paragraph(step, style="List Number")
    doc.add_paragraph("Матрица проверки", "Heading 1")
    table = doc.add_table(rows=1, cols=4)
    table.autofit = False
    for cell, width, title in zip(table.rows[0].cells, (Cm(2.9), Cm(4), Cm(5.1), Cm(5)),
                                  ("Навык", "Материал", "Действие ученика", "Что проверить")):
        cell.width = width
        cell.text = title
        shade(cell, "E8EDF2")
        border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for idx, row in enumerate(ROWS):
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = value
            border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if idx % 2:
                shade(cell, "F7F9FB")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
    doc.add_paragraph("Непроверенное покрытие", "Heading 1")
    doc.add_paragraph("В стартовом вводном сценарии пострадавшего нет. Памятка о новой вводной подготовлена, "
                      "но отдельный стартовый сценарий для проверки этого навыка пока не введён. "
                      "Преподаватель может задать его отдельно; до этого навык нельзя отмечать как пройденный.")
    doc.add_paragraph("Сопоставление с требованиями", "Heading 1")
    doc.add_paragraph("Матрица детализирует раздел об учебных материалах в документе "
                      "«Сопоставление реализации с требованиями ТЗ». Исходный набор билетов содержит 32 билета "
                      "и 96 ситуаций. Их количество не заменяет проверку качества конкретных упражнений.")
    dest = OUT / "14_Матрица_учебных_навыков.docx"
    doc.save(dest)
    return dest


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    for path in (make_materials(), make_coverage()):
        print(path)
