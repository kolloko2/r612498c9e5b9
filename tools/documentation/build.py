"""Build the editable delivery manuals. Run with the bundled document runtime."""
import json
import re
import os
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'output/documentation-2026-09-28'
AUTHOR = 'Кодзаев Николай Петрович'
SIGNATURE = Path(os.environ.get('TRAINER_DOC_SIGNATURE', 'C:/Users/MainUser/Downloads/signature (2).png'))
SCREENS = ROOT / 'docs/delivery/screens'
OUT.mkdir(parents=True, exist_ok=True)
NAMES = ['Ведомость_документов', 'Описание_системы', 'Развёртывание_сервера',
         'Подключение_ПК_ученика', 'Руководство_ученика', 'Руководство_преподавателя',
         'Технологии', 'Соответствие_ТЗ_QA', 'Производительность',
         'Модель_угроз_и_безопасность', 'Программа_испытаний', 'Формуляр', 'Руководство_для_жюри']

# Native editable Word drawings over the original, unmodified screenshots.
# Coordinates refer to the actual 1440 by 960 captured viewport.
CALLOUTS = {
 '02_teacher': [(480,350,640,420,'1'),(1050,190,810,285,'2'),(1070,750,1290,655,'3')],
 '03_lesson': [(600,130,470,270,'1'),(1100,490,960,402,'2'),(600,650,420,558,'3')],
 'teacher_generation': [(1130,220,1050,292,'1'),(1100,620,1050,740,'2'),(750,860,545,877,'3')],
 'teacher_scenarios': [(400,190,450,265,'1'),(730,380,580,325,'2')],
 '01_login': [(1060,330,885,395,'1'),(1060,600,880,550,'2')],
 '04_student': [(530,310,350,350,'1')],
 '05_arm': [(350,390,270,230,'1'),(930,410,860,245,'2'),(410,760,140,915,'3'),(1090,760,1100,920,'4')],
 '06_services': [(380,700,140,880,'1')],
 '06_status': [(1000,680,850,850,'1')],
 '07_tools': [(1130,470,745,700,'1'),(430,460,105,810,'2')],
 '08_report': [(1150,320,800,310,'1'),(1140,470,850,375,'2')],
}
LEGENDS={
 '02_teacher':'1 — создать группу; 2 — открыть сценарии; 3 — назначить отдельное задание.',
 '03_lesson':'1 — выбрать готовые карточки; 2 — задать число карточек; 3 — настроить одновременную работу.',
 'teacher_generation':'1 — режим задания; 2 — описание учебной ситуации; 3 — создание черновика.',
 'teacher_scenarios':'1 — выбор сценария; 2 — переход к генерации с ИИ.',
 '01_login':'1 — учётные данные; 2 — переход после заполнения формы.',
 '04_student':'1 — переход в рабочее место.',
 '05_arm':'1 — адрес; 2 — вид происшествия; 3 — своя служба; 4 — дополнение и инструменты реагирования.',
 '06_services':'1 — блок своей службы и история.',
 '06_status':'1 — панель статуса и комментария.',
 '07_tools':'1 — назначение бригады; 2 — доклад по телефону.',
 '08_report':'1 — итог решений диспетчера; 2 — пояснение невыполненного критерия.',
}

def arrows(p,key):
    scale=459.2126/1440
    for x1,y1,x2,y2,label in CALLOUTS.get(key,[]):
        x1,y1,x2,y2=[v*scale for v in (x1,y1,x2,y2)]
        x1+=9;x2+=9
        drawing_id=max([int(e.get('id')) for e in p.part.element.xpath('.//wp:docPr')]+[0])+1
        sx,sy=(0 if x1<x2 else 1000),(0 if y1<y2 else 1000)
        ex,ey=1000-sx,1000-sy
        emu=lambda v:round(v*12700)
        cx,cy=emu(abs(x2-x1)),emu(abs(y2-y1))
        drawing=f'''<w:drawing xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
        <wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="251659264" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">
        <wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="column"><wp:posOffset>{emu(min(x1,x2))}</wp:posOffset></wp:positionH><wp:positionV relativeFrom="paragraph"><wp:posOffset>{emu(min(y1,y2))}</wp:posOffset></wp:positionV>
        <wp:extent cx="{cx}" cy="{cy}"/><wp:wrapNone/><wp:docPr id="{drawing_id}" name="Указатель {label}"/><wp:cNvGraphicFramePr/>
        <a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:wsp><wps:cNvSpPr/>
        <wps:spPr><a:xfrm flipH="{int(x2<x1)}" flipV="{int(y2<y1)}"><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm><a:prstGeom prst="line"><a:avLst/></a:prstGeom><a:ln w="22860"><a:solidFill><a:srgbClr val="A02020"/></a:solidFill><a:tailEnd type="triangle" w="med" len="med"/></a:ln></wps:spPr><wps:bodyPr/></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing>'''
        p.add_run()._r.append(etree.fromstring(drawing))
        circle=etree.fromstring(drawing)
        ns={'wp':'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing','a':'http://schemas.openxmlformats.org/drawingml/2006/main','wps':'http://schemas.microsoft.com/office/word/2010/wordprocessingShape'}
        circle.find('.//wp:positionH/wp:posOffset',ns).text=str(emu(x1-8))
        circle.find('.//wp:positionV/wp:posOffset',ns).text=str(emu(y1-8))
        for extent in [circle.find('.//wp:extent',ns),circle.find('.//a:xfrm/a:ext',ns)]:
            extent.set('cx',str(emu(16)));extent.set('cy',str(emu(16)))
        circle.find('.//wp:docPr',ns).set('id',str(drawing_id+1))
        circle.find('.//wp:docPr',ns).set('name','Номер '+label)
        circle.find('.//a:prstGeom',ns).set('prst','ellipse')
        xfrm=circle.find('.//a:xfrm',ns);xfrm.set('flipH','0');xfrm.set('flipV','0')
        ln=circle.find('.//a:ln',ns);ln.remove(ln.find('a:tailEnd',ns));ln.set('w','12700')
        sppr=circle.find('.//wps:spPr',ns)
        fill=etree.fromstring('<a:solidFill xmlns:a="'+ns['a']+'"><a:srgbClr val="FFFFFF"/></a:solidFill>');sppr.insert(2,fill)
        wsp=circle.find('.//wps:wsp',ns);body=wsp.find('wps:bodyPr',ns)
        for key,value in {'anchor':'ctr','lIns':'0','rIns':'0','tIns':'0','bIns':'0'}.items():body.set(key,value)
        tx=f'''<wps:txbx xmlns:wps="{ns['wps']}" xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:txbxContent><w:p><w:pPr><w:jc w:val="center"/><w:spacing w:before="0" w:after="0" w:line="220" w:lineRule="exact"/><w:ind w:firstLine="0"/></w:pPr><w:r><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/><w:sz w:val="20"/><w:color w:val="A02020"/></w:rPr><w:t>{label}</w:t></w:r></w:p></w:txbxContent></wps:txbx>'''
        wsp.insert(list(wsp).index(body),etree.fromstring(tx))
        p.add_run()._r.append(circle)

def field(p, name):
    run = p.add_run()
    el = OxmlElement('w:fldSimple'); el.set(qn('w:instr'), name)
    run._r.addnext(el)

def para(doc, text='', style=None):
    p = doc.add_paragraph(text, style)
    return p

def table(doc, rows):
    n = len(rows[0]); t = doc.add_table(rows=0, cols=n)
    t.autofit = False
    weights = [max(8, min(60, sum(len(r[i]) for r in rows)/len(rows))) for i in range(n)]
    if n == 3 and rows[0][0] in ('Код', '№'): weights = [7, 53, 40]
    widths = [16.5*w/sum(weights) for w in weights]
    for i,c in enumerate(t.columns): c.width = Cm(widths[i])
    for ri,row in enumerate(rows):
        cells = t.add_row().cells
        trpr = cells[0]._tc.getparent().get_or_add_trPr()
        no = OxmlElement('w:cantSplit'); trpr.append(no)
        if ri == 0:
            repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
        for ci, text in enumerate(row):
            c = cells[ci]; c.width=Cm(widths[ci]); c.text=text
            pr = c._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side in ('top','left','bottom','right'):
                x=OxmlElement('w:'+side); x.set(qn('w:w'),'80'); x.set(qn('w:type'),'dxa'); margins.append(x)
            pr.append(margins)
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EEEEEE');pr.append(sh)
            for p in c.paragraphs:
                p.paragraph_format.first_line_indent=Cm(0)
                p.paragraph_format.line_spacing=1.05
                p.paragraph_format.space_after=Pt(3)
                p.paragraph_format.keep_with_next=False
                for run in p.runs: run.font.size=Pt(11);run.bold=(ri==0)
    borders=OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
    t._tbl.tblPr.append(borders)
    para(doc).paragraph_format.space_after=Pt(0)

def build(src, idx):
    lines=src.read_text(encoding='utf-8').splitlines()
    jury=idx==12
    meta=dict(line.split(': ',1) for line in lines[:4])
    d=Document(); sec=d.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.left_margin=Cm(3);sec.right_margin=Cm(1.5);sec.top_margin=Cm(2);sec.bottom_margin=Cm(2)
    sec.header_distance=Cm(.8);sec.footer_distance=Cm(1)
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Caption'):
        s=d.styles[name];s.font.name='Times New Roman';s.font.size=Pt(14);s.font.color.rgb=RGBColor(0,0,0)
        s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Times New Roman')
        for attr in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
            s._element.get_or_add_rPr().rFonts.attrib.pop(qn('w:'+attr),None)
        for border in list(s._element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    for border in list(d.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    d.styles['Subtitle'].font.italic=False
    d.styles['Caption'].font.bold=False
    normal=d.styles['Normal'].paragraph_format
    normal.first_line_indent=Cm(1.25);normal.line_spacing=1.5;normal.space_after=Pt(5)
    normal.widow_control=True
    for name in ('Heading 1','Heading 2'):
        p=d.styles[name].paragraph_format;p.first_line_indent=Cm(0);p.space_before=Pt(14);p.space_after=Pt(7);p.keep_with_next=True
        d.styles[name].font.bold=True
    for name in ('Title','Subtitle','Caption'):
        d.styles[name].paragraph_format.first_line_indent=Cm(0)
    d.styles['Title'].font.size=Pt(18)
    d.styles['Caption'].font.size=Pt(12)
    sec.different_first_page_header_footer=True
    hp=sec.header.paragraphs[0];hp.text=meta['Обозначение'];hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.first_line_indent=Cm(0)
    for r in hp.runs:r.font.size=Pt(10)
    fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER;fp.paragraph_format.first_line_indent=Cm(0);field(fp,'PAGE')
    for text,style in [('112 AI Trainer','Title'),('Тренажёр диспетчера ДДС','Subtitle'),(meta['Название'],'Title'),(meta['Обозначение'],None)]:
        p=para(d,text,style);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0)
    para(d)
    for text in ('Программная версия 9fb10f8', 'Разработчик и ответственный за документацию', AUTHOR, '28.09.2026'):
        p=para(d,text);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0)
    p=para(d);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0)
    # Keep the supplied PNG unchanged; crop only its displayed white margins in Word.
    picture=p.add_run().add_picture(str(SIGNATURE),width=Cm(4.4),height=Cm(2.5))
    crop=OxmlElement('a:srcRect')
    for side,value in {'l':'33700','r':'20700','t':'27700','b':'20700'}.items():crop.set(side,value)
    picture._inline.graphic.graphicData.pic.blipFill.insert(1,crop)
    d.add_page_break()
    para(d,'О руководстве' if idx==12 else 'Аннотация','Heading 1');para(d,meta['Аннотация'])
    para(d,'Содержание','Heading 1')
    headings=[];num=0;sub=0
    for line in lines[4:]:
        if line.startswith('## '):num+=1;sub=0;headings.append(line[3:] if jury else f'{num} {line[3:]}')
        elif line.startswith('### ') and not jury:sub+=1;headings.append(f'{num}.{sub} {line[4:]}')
    tocfile=OUT/'toc.json';toc=json.loads(tocfile.read_text(encoding='utf8')) if tocfile.exists() else {}
    for title in headings:
        p=para(d,title+'\t'+str(toc.get(src.stem,{}).get(title,'—')))
        p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(3)
        from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
        p.paragraph_format.tab_stops.add_tab_stop(Cm(16),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
        for r in p.runs:r.font.size=Pt(12)
    d.add_page_break();num=0;sub=0;fig=0;i=4;incode=False
    while i<len(lines):
        line=lines[i];i+=1
        if not line.strip():continue
        if line.startswith('```'):incode=not incode;continue
        if incode:
            p=para(d,line);p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(3)
            for r in p.runs:r.font.name='Courier New';r.font.size=Pt(9)
        elif line.startswith('## '):num+=1;sub=0;para(d,line[3:] if jury else f'{num} {line[3:]}','Heading 1')
        elif line.startswith('### '):sub+=1;para(d,line[4:] if jury else f'{num}.{sub} {line[4:]}','Heading 2')
        elif line.startswith('|'):
            rows=[]
            while True:
                if not re.match(r'^\|[\s:|\-]+\|$',line):rows.append([v.strip() for v in line.strip('|').split('|')])
                if i>=len(lines) or not lines[i].startswith('|'):break
                line=lines[i];i+=1
            table(d,rows)
        elif line.startswith('!figure '):
            key,caption=line[8:].split('|',1);fig+=1
            image=SCREENS/(key+'.png')
            p=para(d);p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(image),width=Cm(16.2))
            arrows(p,key)
            p=para(d,caption if jury else f'Рисунок {fig} — {caption}','Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            if key in LEGENDS:
                p=para(d,LEGENDS[key]);p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1
                for r in p.runs:r.font.size=Pt(11)
        else:para(d,line)
    if idx==12:
        d.part.drop_rel(picture._inline.graphic.graphicData.pic.blipFill.blip.embed)
        # Replace the formal front matter with a plain customer-facing opening.
        content=next(p for p in d.paragraphs if p.text=='Содержание')
        for element in list(d.element.body):
            if element is content._p:break
            d.element.body.remove(element)
        title=content.insert_paragraph_before('112 AI Trainer','Title')
        content.insert_paragraph_before('Как работает тренажёр и как им пользоваться','Subtitle')
        content.insert_paragraph_before(meta['Аннотация'])
        sec.header.paragraphs[0].text=''
        for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Caption'):
            d.styles[name].font.name='Calibri'
        normal.first_line_indent=Cm(0);normal.line_spacing=1.15;normal.space_after=Pt(7)
        d.styles['Normal'].font.size=Pt(12)
        d.styles['Heading 1'].font.size=Pt(17)
        d.styles['Heading 2'].font.size=Pt(14)
        d.styles['Title'].font.size=Pt(22)
    d.core_properties.title=meta['Название'];d.core_properties.subject=meta['Обозначение'];d.core_properties.author=AUTHOR;d.core_properties.language='ru-RU'
    path=OUT/f'{idx:02}_{NAMES[idx]}.docx';d.save(path)
    return {'source':src.stem,'file':path.name,'title':meta['Название'],'headings':headings}

manifest=[build(s,int(s.name[:2])) for s in sorted((ROOT/'docs/delivery/source').glob('*.md')) if not s.name.startswith('11_')]
(OUT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf8')
print(f'Built {len(manifest)} manuals in {OUT}')
