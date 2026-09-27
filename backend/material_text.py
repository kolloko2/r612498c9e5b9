"""Bounded offline document extraction and lexical passage retrieval.

No vector database or network calls. Extract once on save; persist passages in
the material revision, so a lesson does not repeatedly OCR its references.
"""
import hashlib
import io
import json
import os
import re
import shutil
import subprocess
import time
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from text_facts import lemma, tokens

MAX_CHARS = 1_000_000
MAX_UNITS = 200
OCR_SECONDS = 180


def extract(raw, filename, body=''):
    units, warnings = [], []
    started = time.monotonic()
    count = 0
    configured = os.getenv('TESSERACT_CMD', 'tesseract').strip()
    command = json.loads(configured) if configured.startswith('[') else [configured]
    if not isinstance(command, list) or not command or any(not isinstance(x,str) for x in command):
        raise ValueError('TESSERACT_CMD: executable or JSON argument list')
    ocr = shutil.which(command[0])

    def add(text, location):
        nonlocal count
        text = re.sub(r'[ \t]+', ' ', text or '').strip()
        remaining = MAX_CHARS-count
        if len(text) > remaining:
            warnings.append('Достигнут предел извлечённого текста: 1 000 000 символов.')
        text = text[:max(0, remaining)]
        count += len(text)
        for offset in range(0, len(text), 1600):
            fragment = text[offset:offset+1800]
            units.append({'location':location, 'text':fragment,
                          '_terms':sorted({lemma(w) for w in tokens(fragment) if len(w)>2})})

    def recognize(image, location):
        if not ocr:
            warnings.append('Для изображений нужен локальный Tesseract с языками rus и eng.')
            return
        remaining = OCR_SECONDS-(time.monotonic()-started)
        if remaining <= 0:
            warnings.append('Истёк лимит OCR 180 секунд; часть изображений не прочитана.')
            return
        try:
            stream = io.BytesIO()
            image.save(stream, format='PNG')
            result = subprocess.run([ocr, *command[1:], 'stdin', 'stdout', '-l', 'rus+eng'],
                input=stream.getvalue(), stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                timeout=min(30, remaining), check=True)
            text = result.stdout.decode('utf-8').strip()
            if not text:
                warnings.append(location + ': текст на изображении не распознан.')
            add(text, location + ' (OCR)')
        except (subprocess.SubprocessError, OSError, UnicodeError):
            warnings.append(location + ': OCR не завершён; проверьте локальный движок и языки.')

    add(body, 'Текст преподавателя')
    pages = 0
    if raw:
        try:
            ext = filename.lower().rsplit('.', 1)[-1]
            if ext == 'txt':
                add(raw.decode('utf-8-sig'), 'Файл')
            elif ext == 'pdf':
                import pypdfium2 as pdfium
                pdf = pdfium.PdfDocument(raw)
                try:
                    pages = len(pdf)
                    if pages > MAX_UNITS:
                        warnings.append('Обработаны первые 200 страниц; разделите документ.')
                    for i in range(min(pages, MAX_UNITS)):
                        page = pdf[i]
                        try:
                            textpage = page.get_textpage()
                            try:
                                text = textpage.get_text_range()
                            finally:
                                textpage.close()
                            if len(re.sub(r'\W', '', text)) >= 30:
                                add(text, f'Страница {i+1}')
                            else:
                                # Cap pixels as well as page count for untrusted PDFs.
                                scale = min(2, (12_000_000/max(1, page.get_width()*page.get_height()))**.5)
                                bitmap = page.render(scale=scale)
                                try:
                                    recognize(bitmap.to_pil(), f'Страница {i+1}')
                                finally:
                                    bitmap.close()
                        finally:
                            page.close()
                finally:
                    pdf.close()
            elif ext == 'docx':
                from PIL import Image
                with ZipFile(io.BytesIO(raw)) as archive:
                    root = ET.fromstring(archive.read('word/document.xml'))
                    paragraphs = [' '.join(n.text or '' for n in p.iter() if n.tag.endswith('}t'))
                                  for p in root.iter() if p.tag.endswith('}p')]
                    add('\n'.join(paragraphs), 'Документ')
                    seen = set()
                    images = [n for n in archive.namelist() if n.startswith('word/media/')]
                    if len(images) > MAX_UNITS:
                        warnings.append('Обработаны первые 200 изображений; разделите документ.')
                    for name in images[:MAX_UNITS]:
                        data = archive.read(name)
                        digest = hashlib.sha256(data).digest()
                        if digest in seen:
                            continue
                        seen.add(digest)
                        try:
                            with Image.open(io.BytesIO(data)) as image:
                                if image.width*image.height > 25_000_000:
                                    warnings.append(name+': изображение слишком велико для OCR.')
                                else:
                                    recognize(image, name)
                        except (OSError, ValueError):
                            warnings.append(name+': формат изображения не распознан.')
                    pages = len(images)
            elif ext == 'xlsx':
                import openpyxl
                workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
                try:
                    for sheet in workbook:
                        if sheet.max_row > 10000 or sheet.max_column > 200:
                            warnings.append(sheet.title+': прочитан диапазон до 10000 строк и 200 столбцов.')
                        for i, row in enumerate(sheet.iter_rows(max_row=min(sheet.max_row, 10000), max_col=min(sheet.max_column, 200), values_only=True), 1):
                            text = ' | '.join(str(v) if v is not None else '' for v in row)
                            if text.strip(' |'):
                                add(text, f'{sheet.title}, строка {i}')
                            if count >= MAX_CHARS:
                                break
                finally:
                    workbook.close()
        except Exception:
            # Failed parsing is visible, and never turns into a claim of full ingestion.
            warnings.append('Не удалось полностью прочитать вложение. Проверьте файл и извлечённый текст.')
    warnings = list(dict.fromkeys(warnings))
    return units, {'status':'partial' if warnings else ('ready' if units else 'empty'),
                   'characters':count, 'passages':len(units), 'pages_or_images':pages,
                   'warnings':warnings, 'preview':'\n\n'.join(u['location']+'\n'+u['text'] for u in units)[:6000]}


def retrieve(documents, query='', budget=12000):
    """Rank all approved passages by current task, retaining source locations."""
    wanted = {lemma(w) for w in tokens(query) if len(w) > 2}
    candidates = []
    for document in documents:
        for index, chunk in enumerate(document.get('_passages', [])):
            words = set(chunk['_terms']) if '_terms' in chunk else {lemma(w) for w in tokens(chunk['text']) if len(w) > 2}
            score = len(words & wanted) / max(1, len(wanted))
            candidates.append((score, document['updated_at'], -index, document, chunk))
    candidates.sort(key=lambda row:row[:3], reverse=True)
    result, used = [], 0
    for _, _, _, document, chunk in candidates:
        if used >= budget:
            break
        text = chunk['text'][:budget-used]
        result.append({'id':document['id'], 'title':document['title'], 'revision':document['revision'],
                       'location':chunk['location'], 'excerpt':text})
        used += len(text)
    return result
