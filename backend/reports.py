"""XLSX export of an already authorized statistics snapshot.

The workbook is built from exactly the dictionary the statistics endpoint returns,
so a teacher can never export a row the API would not have shown them. Nothing is
recomputed here and no extra query is issued.

Cells are written as text or numbers only. A string that a spreadsheet could read
as a formula is prefixed with an apostrophe, matching the existing CSV export rule
in docs/ASSESSMENT.md: an exported training report must never execute on opening.
"""

from __future__ import annotations

import io
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

FORMULA_MARKERS = ("=", "+", "-", "@", "\t", "\r")
MAX_SHEET_TITLE = 31
# Светло-зелёный → красный. Шкала только визуальная: она повторяет rate_percent
# и не вводит собственного порога «допустимой» доли ошибок.
HEAT_STEPS = ((0, "C8E6C9"), (20, "DCEDC8"), (40, "FFF9C4"), (60, "FFE0B2"), (80, "FFCDD2"))

SUMMARY_COLUMNS = ("attempts", "completed", "graded", "passed", "failed",
                   "unassessed", "average_score", "average_seconds")


def _safe(value: Any) -> Any:
    """Never hand a spreadsheet something it can evaluate."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (int, float)):
        return value
    text = str(value)
    if text.lstrip().startswith(FORMULA_MARKERS):
        return "'" + text
    return text


def _sheet(workbook: Workbook, title: str, headers: Iterable[str], rows: Iterable[Iterable[Any]]):
    sheet = workbook.create_sheet(title[:MAX_SHEET_TITLE])
    headers = list(headers)
    sheet.append([_safe(header) for header in headers])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([_safe(value) for value in row])
    for index, header in enumerate(headers, start=1):
        width = max(len(str(header)) + 2, 12)
        sheet.column_dimensions[get_column_letter(index)].width = min(width, 46)
    sheet.freeze_panes = "A2"
    return sheet


def _fill(rate: float) -> PatternFill:
    colour = HEAT_STEPS[0][1]
    for threshold, value in HEAT_STEPS:
        if rate >= threshold:
            colour = value
    return PatternFill("solid", fgColor=colour)


def workbook_bytes(statistics: dict, *, scope: str) -> bytes:
    """Render one statistics snapshot; `scope` names the exporting role in a note."""
    workbook = Workbook()
    workbook.remove(workbook.active)

    summary = statistics.get("summary", {})
    _sheet(workbook, "Сводка", ("Показатель", "Значение"),
           [(column, summary.get(column)) for column in SUMMARY_COLUMNS])

    progress = statistics.get("progress", [])
    progress_columns = ("session_id", "title", "finished_at", "student_id", "student_name",
                        "score_percent", "passed", "source", "difficulty", "dds_profile")
    _sheet(workbook, "Прогресс", progress_columns,
           [[row.get(column) for column in progress_columns] for row in progress])

    students = statistics.get("students", [])
    if students:
        student_columns = ("student_id", "display_name", *SUMMARY_COLUMNS)
        _sheet(workbook, "Обучающиеся", student_columns,
               [[row.get(column) for column in student_columns] for row in students])

    errors = statistics.get("typical_errors", [])
    error_columns = ("key", "label", "count", "attempts", "rate_percent")
    _sheet(workbook, "Типичные ошибки", error_columns,
           [[row.get(column) for column in error_columns] for row in errors])

    heatmap = statistics.get("error_heatmap") or {}
    cells = heatmap.get("cells", [])
    checks = heatmap.get("checks", [])
    scenarios = heatmap.get("scenarios", [])
    index = {(cell["scenario_id"], cell["label"]): cell for cell in cells}
    sheet = _sheet(workbook, "Тепловая карта", ("Сценарий", *checks), [])
    for scenario in scenarios:
        row = [scenario.get("title") or scenario["scenario_id"]]
        for check in checks:
            cell = index.get((scenario["scenario_id"], check))
            row.append(cell["rate_percent"] if cell else None)
        sheet.append([_safe(value) for value in row])
        for column, check in enumerate(checks, start=2):
            cell = index.get((scenario["scenario_id"], check))
            if cell:
                target = sheet.cell(row=sheet.max_row, column=column)
                target.fill = _fill(cell["rate_percent"])
                target.number_format = "0.00"
    notes = workbook.create_sheet("О выгрузке")
    notes.append(["Выгрузка учебной статистики тренажёра 112"])
    notes.append(["Область выгрузки", scope])
    notes.append(["Доля ошибок, %", "значение ячейки тепловой карты = rate_percent"])
    notes.column_dimensions["A"].width = 26
    notes.column_dimensions["B"].width = 70
    notes["A1"].font = Font(bold=True)

    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()
