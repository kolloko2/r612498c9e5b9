import io

from openpyxl import load_workbook

from reports import workbook_bytes

SNAPSHOT = {
    'summary': {'attempts': 3, 'completed': 2, 'graded': 2, 'passed': 1, 'failed': 1,
                'unassessed': 0, 'average_score': 50, 'average_seconds': 12},
    'progress': [{'session_id': 's1', 'title': '=HYPERLINK("http://evil","click")',
                  'finished_at': '2026-09-15T10:00:00+00:00', 'student_id': 'u1',
                  'student_name': '@student', 'score_percent': 50, 'passed': False,
                  'source': 'automatic', 'difficulty': 'basic', 'dds_profile': 'fire'}],
    'students': [{'student_id': 'u1', 'display_name': '-Иванов', 'attempts': 3, 'completed': 2,
                  'graded': 2, 'passed': 1, 'failed': 1, 'unassessed': 0,
                  'average_score': 50, 'average_seconds': 12}],
    'typical_errors': [{'key': 'field:t:sc:1:caller', 'label': 'Заявитель',
                        'count': 1, 'attempts': 2, 'rate_percent': 50.0}],
    'error_heatmap': {
        'scenarios': [{'scenario_id': 'sc', 'title': 'Пожар в квартире'}],
        'checks': ['Заявитель', 'Улица'],
        'cells': [
            {'scenario_id': 'sc', 'title': 'Пожар в квартире', 'label': 'Заявитель',
             'count': 1, 'attempts': 2, 'rate_percent': 50.0},
            {'scenario_id': 'sc', 'title': 'Пожар в квартире', 'label': 'Улица',
             'count': 0, 'attempts': 2, 'rate_percent': 0.0},
        ],
    },
}


def book(snapshot=SNAPSHOT, scope='тест'):
    return load_workbook(io.BytesIO(workbook_bytes(snapshot, scope=scope)))


def test_workbook_sheets_and_summary():
    workbook = book()
    assert workbook.sheetnames == ['Сводка', 'Прогресс', 'Обучающиеся',
                                   'Типичные ошибки', 'Тепловая карта', 'О выгрузке']
    rows = {row[0]: row[1] for row in workbook['Сводка'].iter_rows(min_row=2, values_only=True)}
    assert rows['attempts'] == 3 and rows['average_score'] == 50


def test_formula_like_text_cannot_execute():
    """An exported report must never run a formula when a reviewer opens it."""
    workbook = book()
    title = workbook['Прогресс']['B2'].value
    name = workbook['Прогресс']['E2'].value
    student = workbook['Обучающиеся']['B2'].value
    assert title.startswith("'=") and name.startswith("'@") and student.startswith("'-")
    for sheet in workbook.sheetnames:
        for row in workbook[sheet].iter_rows(values_only=True):
            for value in row:
                if isinstance(value, str):
                    assert not value.startswith(('=', '+', '@', '\t', '\r'))


def test_heatmap_cells_carry_rate_and_colour():
    workbook = book()
    sheet = workbook['Тепловая карта']
    assert [cell.value for cell in sheet[1]] == ['Сценарий', 'Заявитель', 'Улица']
    assert sheet['A2'].value == 'Пожар в квартире'
    assert sheet['B2'].value == 50.0 and sheet['C2'].value == 0.0
    # Разная доля ошибок — разная заливка, иначе карта ничего не показывает.
    assert sheet['B2'].fill.fgColor.rgb != sheet['C2'].fill.fgColor.rgb


def test_empty_snapshot_still_produces_a_readable_workbook():
    workbook = book({'summary': {}, 'progress': [], 'typical_errors': []})
    assert 'Обучающиеся' not in workbook.sheetnames
    assert workbook['Тепловая карта'].max_row == 1
    assert workbook['О выгрузке']['B2'].value == 'тест'
