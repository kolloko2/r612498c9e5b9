import json
from uuid import uuid4

import pytest
from pydantic import ValidationError
from accounts import Accounts
from learning import Learning
from assessment import router, AssessmentPolicy, evaluate_policy
from test_rbac_integration import classroom
from test_lesson_lifecycle import create_lesson


@pytest.fixture
def grading(classroom):
    c = classroom
    accounts = Accounts(c['store'])
    c['client'].app.include_router(router(c['store'], accounts, Learning(c['store'], accounts), lambda: None))
    return c


def policy_url(c):
    return '/api/v1/instructor/scenarios/'+c['scenario_id']+'/assessment-policy'


def policy(**changes):
    return {'pass_score_percent': 80, 'max_field_errors': 0, 'max_sequence_errors': 0,
            'steps': [{'id': 'save', 'label': 'Сохранить карточку', 'event_type': 'card.saved'},
                      {'id': 'finish', 'label': 'Завершить', 'event_type': 'session.finished'}], **changes}


def configure(c):
    client = c['client']; h = c['headers']['teacher1']
    assert client.put(policy_url(c), headers=h, json={'revision': 0, 'policy': policy()}).status_code == 200
    assert client.put('/api/v1/instructor/scenarios/'+c['scenario_id']+'/rubric', headers=h, json={'revision': 0, 'rubric': {
        'title': 'Учебный эталон', 'time_limit_seconds': 120, 'criteria': [
            {'id': 'name', 'label': 'Заявитель', 'field': 'caller_name', 'mode': 'equals', 'expected': ['Учебный'], 'weight': 1}]}}).status_code == 200


def new_card(c):
    result = c['client'].post('/api/v1/student/sessions', headers=c['headers']['student1'], json={
        'scenario_id': c['scenario_id'], 'assignment_id': c['assignment']['id']})
    assert result.status_code == 201, result.text
    return result.json()


def finish(c, card):
    result = c['client'].post('/api/v1/student/sessions/'+card['id']+'/finish', headers=c['headers']['student1'])
    assert result.status_code == 200, result.text
    return result.json()


def test_policy_freeze_errors_and_completed_immutability(grading):
    c = grading; configure(c); card = new_card(c); h = c['headers']['teacher1']
    assert 'assessment_policy' not in card and 'policy_result' not in card
    assert c['client'].put(policy_url(c), headers=h, json={'revision': 1, 'policy': None}).status_code == 200
    result = finish(c, card)
    assert result['evaluation']['score_percent'] == 0
    assert result['policy_result']['passed'] is False
    assert result['policy_result']['field_errors'] == 1
    assert result['policy_result']['sequence_errors'] == 1
    assert result['policy_result']['steps'][-1]['event_seq'] == result['events'][-1]['seq']
    assert result['policy_result']['policy_revision'] == 1
    assert finish(c, card) == result
    assert finish(c, new_card(c))['policy_result'] is None


def test_policy_role_conflict_and_history(grading):
    c = grading; client = c['client']; h = c['headers']; url = policy_url(c)
    assert client.put(url, headers=h['student1'], json={'revision': 0, 'policy': policy()}).status_code == 403
    assert client.put(url, headers=h['teacher1'], json={'revision': 0, 'policy': policy()}).status_code == 200
    assert client.put(url, headers=h['teacher1'], json={'revision': 0, 'policy': None}).status_code == 409
    assert client.get(url, headers=h['teacher2']).json() == {'revision': 0, 'policy': None}
    assert client.put(url, headers=h['teacher1'], json={'revision': 1, 'policy': {}}).status_code == 422
    assert c['store'].db.execute('SELECT count(*) FROM assessment_policy_history').fetchone()[0] == 1


def test_order_repeated_steps_filters_and_missing_score():
    p = {'revision': 2, 'policy': policy(pass_score_percent=None, max_field_errors=None, steps=[
        {'id': 'save', 'label': 'Сохранить', 'event_type': 'card.saved'},
        {'id': 'accepted', 'label': 'Принять', 'event_type': 'service.updated', 'service': 'Служба 101', 'status': 'Принята'},
        {'id': 'save2', 'label': 'Сохранить ещё', 'event_type': 'card.saved'}])}
    events = [{'seq': 1, 'type': 'service.updated', 'detail': {'service': 'Служба 101', 'status': 'Принята'}}, {'seq': 2, 'type': 'card.saved'}]
    result = evaluate_policy(p, {'status': 'not_configured'}, events)
    assert result['sequence_errors'] == 2 and result['steps'][1]['reason'] == 'out_of_order'
    events.extend([{'seq': 3, 'type': 'service.updated', 'detail': {'service': 'Служба 101', 'status': 'Принята'}}, {'seq': 4, 'type': 'card.saved'}])
    assert evaluate_policy(p, {'status': 'not_configured'}, events)['passed'] is True
    unknown = evaluate_policy({'revision': 1, 'policy': policy(steps=[])}, {'status': 'not_configured'}, [])
    assert unknown['passed'] is None and unknown['field_errors'] is None
    timed = evaluate_policy({'revision': 1, 'policy': {'fail_on_timeout': True}}, {'timing': {'within_limit': False, 'elapsed_seconds': 31, 'limit_seconds': 30}}, [])
    assert timed['passed'] is False
    with pytest.raises(ValidationError):
        AssessmentPolicy(steps=[{'id': 'x', 'label': 'x', 'event_type': 'card.saved', 'service': 'Invalid'}])


def test_expert_append_retry_revoke_and_privacy(grading):
    c = grading; configure(c); card = new_card(c); client = c['client']; h = c['headers']
    url = '/api/v1/instructor/sessions/'+card['id']+'/assessment'
    request = {'request_id': str(uuid4()), 'revision': 0, 'score_percent': 90, 'passed': True, 'reason': 'Проверены обстоятельства учебной карточки'}
    assert client.post(url, headers=h['teacher1'], json=request).status_code == 409
    original = finish(c, card)
    response = client.post(url, headers=h['teacher1'], json=request)
    assert response.status_code == 200, response.text
    result = response.json()
    assert result['effective'] == {'score_percent': 90, 'passed': True, 'source': 'expert'}
    assert result['automatic'] == original['evaluation']
    assert client.post(url, headers=h['teacher1'], json=request).json() == result
    assert client.post(url, headers=h['teacher1'], json={**request, 'score_percent': 89}).status_code == 409
    assert client.post(url, headers=h['teacher1'], json={**request, 'request_id': str(uuid4())}).status_code == 409
    assert client.get(url, headers=h['teacher2']).status_code == 404
    assert client.post(url, headers=h['student1'], json=request).status_code == 403
    student_url = '/api/v1/student/sessions/'+card['id']+'/assessment'
    assert client.get(student_url, headers=h['student2']).status_code == 404
    assert client.get(student_url, headers=h['student1']).json()['expert']['current']['teacher_name'] == 'Преподаватель 1'
    revoked = client.post(url, headers=h['teacher1'], json={'request_id': str(uuid4()), 'revision': 1, 'action': 'revoke', 'reason': 'Возврат к автоматическому результату'}).json()
    assert revoked['expert']['current'] is None and len(revoked['expert']['history']) == 2
    assert revoked['effective']['score_percent'] == 0 and revoked['effective']['passed'] is False
    persisted = client.get('/api/v1/student/sessions/'+card['id'], headers=h['student1']).json()
    assert persisted == original


def test_statistics_effective_grade_error_denominators_and_group_privacy(grading):
    c = grading; configure(c); client = c['client']; h = c['headers']
    failed = finish(c, new_card(c)); card = new_card(c)
    assert client.put('/api/v1/student/sessions/'+card['id']+'/card', headers=h['student1'], json={'revision': 0, 'card': {**card['card'], 'caller_name': 'Учебный'}}).status_code == 200
    success = finish(c, card)
    assert success['policy_result']['passed'] is True
    stats = client.get('/api/v1/student/statistics', headers=h['student1']).json()
    assert stats['summary'] == {'attempts': 3, 'completed': 2, 'graded': 2, 'passed': 1, 'failed': 1, 'unassessed': 0, 'average_score': 50, 'average_seconds': stats['summary']['average_seconds']}
    assert stats['students'] == []
    name_error = next(r for r in stats['typical_errors'] if r['label'] == 'Заявитель')
    assert name_error['count'] == 1 and name_error['attempts'] == 2 and name_error['rate_percent'] == 50
    request = {'request_id': str(uuid4()), 'revision': 0, 'score_percent': 80, 'passed': True, 'reason': 'Экспертный пересмотр'}
    assert client.post('/api/v1/instructor/sessions/'+failed['id']+'/assessment', headers=h['teacher1'], json=request).status_code == 200
    updated = client.get('/api/v1/student/statistics', headers=h['student1']).json()
    assert updated['summary']['average_score'] == 90 and updated['summary']['passed'] == 2
    assert updated['typical_errors'] == stats['typical_errors']
    assert client.get('/api/v1/student/statistics', headers=h['student2']).json()['summary']['attempts'] == 0
    group = c['group']['id']
    assert client.get('/api/v1/instructor/statistics?group_id='+group, headers=h['teacher2']).status_code == 404
    assert client.post('/api/v1/instructor/groups/'+group+'/members', headers=h['teacher1'], json={'student_id': c['users']['student2']['id']}).status_code == 200
    grouped = client.get('/api/v1/instructor/statistics?group_id='+group, headers=h['teacher1']).json()
    assert len(grouped['students']) == 2 and grouped['students'][1]['completed'] in (0, 2)
    assert sum(s['completed'] for s in grouped['students']) == 2


def test_lesson_freezes_policy_and_action_only_sequence(grading):
    c = grading; configure(c); client = c['client']; h = c['headers']
    lesson = create_lesson(c).json()
    assert client.post('/api/v1/instructor/lessons/'+lesson['id']+'/start', headers=h['teacher1']).status_code == 200
    action_policy = policy(pass_score_percent=None, max_field_errors=None, steps=[{'id': 'finish', 'label': 'Завершить', 'event_type': 'session.finished'}])
    assert client.put(policy_url(c), headers=h['teacher1'], json={'revision': 1, 'policy': action_policy}).status_code == 200
    issued = client.post('/api/v1/student/lessons/'+lesson['id']+'/next', headers=h['student1']).json()
    assert finish(c, issued)['policy_result']['policy_revision'] == 1
    actions = create_lesson(c, scenario_ids=[], mode='actions', prefilled_scenario_ids=[c['scenario_id']]).json()
    assert client.put(policy_url(c), headers=h['teacher1'], json={'revision': 2, 'policy': None}).status_code == 200
    assert client.post('/api/v1/instructor/lessons/'+actions['id']+'/start', headers=h['teacher1']).status_code == 200
    issued = client.post('/api/v1/student/lessons/'+actions['id']+'/next', headers=h['student1']).json()
    result = finish(c, issued)
    assert result['evaluation']['score_percent'] is None and result['policy_result']['passed'] is True
    assert result['policy_result']['policy_revision'] == 2


def test_error_heatmap_matches_typical_errors(grading):
    c = grading; configure(c); client = c['client']; h = c['headers']
    finish(c, new_card(c)); card = new_card(c)
    client.put('/api/v1/student/sessions/'+card['id']+'/card', headers=h['student1'],
               json={'revision': 0, 'card': {**card['card'], 'caller_name': 'Учебный'}})
    finish(c, card)
    stats = client.get('/api/v1/student/statistics', headers=h['student1']).json()
    heatmap = stats['error_heatmap']
    assert heatmap['scenarios'] and heatmap['checks']
    # Тепловая карта пересчитывает те же исходы, что и список типичных ошибок.
    name_cells = [cell for cell in heatmap['cells'] if cell['label'] == 'Заявитель']
    name_error = next(r for r in stats['typical_errors'] if r['label'] == 'Заявитель')
    assert sum(cell['count'] for cell in name_cells) == name_error['count']
    assert sum(cell['attempts'] for cell in name_cells) == name_error['attempts']
    assert all(0 <= cell['rate_percent'] <= 100 for cell in heatmap['cells'])
    # Чужая статистика остаётся пустой и в тепловой карте.
    assert client.get('/api/v1/student/statistics', headers=h['student2']).json()['error_heatmap']['cells'] == []


def test_statistics_workbook_export(grading):
    import base64, io
    from openpyxl import load_workbook
    c = grading; configure(c); client = c['client']; h = c['headers']
    finish(c, new_card(c))

    def workbook(payload):
        return load_workbook(io.BytesIO(base64.b64decode(payload['file_base64'])))

    response = client.get('/api/v1/student/statistics-workbook', headers=h['student1'])
    assert response.status_code == 200
    payload = response.json()
    assert payload['filename'].endswith('.xlsx')
    assert payload['content_type'].endswith('spreadsheetml.sheet')
    book = workbook(payload)
    assert {'Сводка', 'Прогресс', 'Типичные ошибки', 'Тепловая карта', 'О выгрузке'} <= set(book.sheetnames)
    assert book['Сводка']['A2'].value == 'attempts'
    # Обучающийся не получает лист с другими людьми.
    assert 'Обучающиеся' not in book.sheetnames
    teacher = client.get('/api/v1/instructor/statistics-workbook', headers=h['teacher1'])
    assert teacher.status_code == 200
    assert 'Обучающиеся' in workbook(teacher.json()).sheetnames
    # Роли не пересекаются.
    assert client.get('/api/v1/instructor/statistics-workbook', headers=h['student1']).status_code == 403
    assert client.get('/api/v1/student/statistics-workbook', headers=h['teacher1']).status_code == 403


def test_same_check_in_different_scenarios_is_one_typical_error(grading):
    c = grading; client = c['client']; h = c['headers']
    student = c['users']['student1']['id']; teacher = c['users']['teacher1']['id']
    rows = []
    for index, scenario in enumerate(('scenario-a', 'scenario-b')):
        rows.append({'id': f'dds-{index}', 'status': 'Завершена', 'student_id': student, 'teacher_id': teacher,
                     'scenario_id': scenario, 'title': scenario, 'exercise_mode': 'actions',
                     'created_at': '2026-09-28T10:00:00+00:00', 'finished_at': '2026-09-28T10:05:00+00:00',
                     'card': {}, 'events': [],
                     'dds_review': {'version': f'v{index}', 'score_percent': 50, 'passed': False, 'checks': [
                         {'id': 'briefing', 'label': 'Доклад вышестоящему начальнику', 'passed': False}]},
                     # Поля готовой карточки ДДС ученик не заполнял: это не его ошибка.
                     'evaluation': {'status': 'evaluated', 'rubric_revision': 1, 'criteria': [
                         {'id': 'name', 'label': 'ФИО заявителя', 'passed': False}]}})
    with c['store'].db:
        for row in rows:
            c['store'].db.execute('INSERT INTO workspace VALUES (?,?)', (row['id'], json.dumps(row, ensure_ascii=False)))
    stats = client.get('/api/v1/student/statistics', headers=h['student1']).json()
    labels = [item['label'] for item in stats['typical_errors']]
    assert labels.count('Доклад вышестоящему начальнику') == 1
    assert next(i for i in stats['typical_errors'] if i['label'] == 'Доклад вышестоящему начальнику')['count'] == 2
    assert 'ФИО заявителя' not in labels
